from types import SimpleNamespace
from io import BytesIO
from datetime import date

from django.test import TestCase
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from location.gql_mutations import update_or_create_hotspot
from location.micro_catchment_workbook import build_template_workbook, build_workbook, import_excel, import_records
from location.models import Hotspot, Location, MicroCatchment, MicroCatchmentTA
from location.services import MicroCatchmentService


class HierarchicalCodeGenerationTest(TestCase):
    def setUp(self):
        self.user = SimpleNamespace(
            id_for_audit=-1,
            is_superuser=True,
            is_anonymous=False,
        )
        self.district = self._location("R", "DIST")
        self.ta = self._location("D", "2341", self.district)
        self.gvh = self._location("W", "GVH1", self.ta)

    def _location(self, location_type, code, parent=None):
        return Location.objects.create(
            type=location_type,
            code=code,
            name=code,
            parent=parent,
            audit_user_id=-1,
        )

    def _create_micro_catchment(self, name, gvh=None):
        gvh = gvh or self.gvh
        return MicroCatchmentService(self.user).update_or_create(
            {
                "name": name,
                "district_id": self.district.id,
                "ta_ids": [self.ta.id],
                "gvh_ids": [gvh.id],
                "audit_user_id": -1,
            }
        )

    def test_micro_catchment_code_uses_ta_prefix_and_increments(self):
        first = self._create_micro_catchment("First")
        second = self._create_micro_catchment("Second")

        self.assertEqual(first.code, "234101")
        self.assertEqual(second.code, "234102")

    def test_micro_catchment_suffix_expands_beyond_two_digits(self):
        existing = MicroCatchment.objects.create(
            code="234199",
            name="Existing",
            district=self.district,
            audit_user_id=-1,
        )
        MicroCatchmentTA.objects.create(
            micro_catchment=existing,
            location=self.ta,
            audit_user_id=-1,
        )

        created = self._create_micro_catchment("One hundred")

        self.assertEqual(created.code, "2341100")

    def test_hotspot_code_uses_micro_catchment_prefix_and_increments(self):
        micro_catchment = self._create_micro_catchment("Catchment")
        first_village = self._location("V", "V1", self.gvh)
        second_village = self._location("V", "V2", self.gvh)

        first = update_or_create_hotspot(
            {
                "name": "First",
                "micro_catchment_uuid": micro_catchment.uuid,
                "village_uuids": [first_village.uuid],
                "audit_user_id": -1,
            },
            self.user,
        )
        second = update_or_create_hotspot(
            {
                "name": "Second",
                "micro_catchment_uuid": micro_catchment.uuid,
                "village_uuids": [second_village.uuid],
                "audit_user_id": -1,
            },
            self.user,
        )

        self.assertEqual(first.code, f"{micro_catchment.code}01")
        self.assertEqual(second.code, f"{micro_catchment.code}02")

    def test_hotspot_suffix_expands_beyond_two_digits(self):
        micro_catchment = self._create_micro_catchment("Catchment")
        Hotspot.objects.create(
            code=f"{micro_catchment.code}99",
            name="Existing",
            micro_catchment=micro_catchment,
            audit_user_id=-1,
        )
        village = self._location("V", "V100", self.gvh)

        created = update_or_create_hotspot(
            {
                "name": "One hundred",
                "micro_catchment_uuid": micro_catchment.uuid,
                "village_uuids": [village.uuid],
                "audit_user_id": -1,
            },
            self.user,
        )

        self.assertEqual(created.code, f"{micro_catchment.code}100")

    def test_moved_hotspot_code_is_not_reused_in_original_catchment(self):
        first_catchment = self._create_micro_catchment("First catchment")
        second_catchment = self._create_micro_catchment("Second catchment")
        first_village = self._location("V", "MOVE-V1", self.gvh)
        second_village = self._location("V", "MOVE-V2", self.gvh)
        moved_hotspot = update_or_create_hotspot(
            {
                "name": "Moved hotspot",
                "micro_catchment_uuid": first_catchment.uuid,
                "village_uuids": [first_village.uuid],
                "audit_user_id": -1,
            },
            self.user,
        )
        original_code = moved_hotspot.code

        update_or_create_hotspot(
            {
                "uuid": moved_hotspot.uuid,
                "name": moved_hotspot.name,
                "micro_catchment_uuid": second_catchment.uuid,
                "village_uuids": [first_village.uuid],
                "audit_user_id": -1,
            },
            self.user,
        )
        replacement = update_or_create_hotspot(
            {
                "name": "Replacement hotspot",
                "micro_catchment_uuid": first_catchment.uuid,
                "village_uuids": [second_village.uuid],
                "audit_user_id": -1,
            },
            self.user,
        )

        self.assertEqual(original_code, f"{first_catchment.code}01")
        self.assertEqual(replacement.code, f"{first_catchment.code}02")

    def test_import_template_omits_micro_catchment_code(self):
        workbook = load_workbook(BytesIO(build_template_workbook(self.district)))
        headers = [cell.value for cell in workbook["MicroCatchments"][1]]

        self.assertNotIn("micro_catchment_code", headers)
        self.assertIn("start_date", headers)
        self.assertIn("end_date", headers)
        self.assertIn("micro_catchment_name", headers)

    def test_export_includes_micro_catchment_dates(self):
        catchment = self._create_micro_catchment("Dated catchment")
        catchment.date_from = date(2026, 8, 1)
        catchment.date_to = date(2026, 8, 31)
        catchment.save(update_fields=["date_from", "date_to"])
        workbook = load_workbook(BytesIO(build_workbook(self.district)))
        sheet = workbook["MicroCatchments"]
        headers = [cell.value for cell in sheet[1]]
        values = dict(zip(headers, [cell.value for cell in sheet[2]]))

        self.assertIn("date_from", headers)
        self.assertIn("date_to", headers)
        self.assertEqual(str(values["date_from"])[:10], "2026-08-01")
        self.assertEqual(str(values["date_to"])[:10], "2026-08-31")

    def test_excel_import_generates_code_from_ta(self):
        content = build_template_workbook(self.district)
        workbook = load_workbook(BytesIO(content))
        sheet = workbook["MicroCatchments"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, headers["micro_catchment_name"], "Imported catchment")
        sheet.cell(2, headers["start_date"], "2026-08-01")
        sheet.cell(2, headers["end_date"], "2026-08-31")
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        result = import_excel(uploaded, self.district, self.user)

        imported = MicroCatchment.objects.get(name="Imported catchment")
        self.assertEqual(result, {"created": 1, "updated": 0, "total": 1})
        self.assertEqual(imported.code, "234101")
        self.assertEqual(imported.date_from, date(2026, 8, 1))
        self.assertEqual(imported.date_to, date(2026, 8, 31))

    def test_excel_import_rejects_end_date_before_start_date(self):
        content = build_template_workbook(self.district)
        workbook = load_workbook(BytesIO(content))
        sheet = workbook["MicroCatchments"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, headers["micro_catchment_name"], "Invalid dates")
        sheet.cell(2, headers["start_date"], "2026-08-31")
        sheet.cell(2, headers["end_date"], "2026-08-01")
        uploaded = BytesIO()
        workbook.save(uploaded)
        uploaded.seek(0)

        with self.assertRaises(ValidationError) as context:
            import_excel(uploaded, self.district, self.user)

        self.assertIn("end_date cannot be before start_date", " ".join(context.exception.messages))

    def test_import_rejects_existing_micro_catchment_name(self):
        self._create_micro_catchment("Existing catchment")

        with self.assertRaises(ValidationError) as context:
            import_records(
                [
                    {
                        "name": "existing CATCHMENT",
                        "type": None,
                        "tas": [self.ta],
                        "gvhs": [self.gvh],
                    }
                ],
                self.district,
                self.user,
            )

        self.assertIn("already exists", " ".join(context.exception.messages))

    def test_service_rejects_duplicate_micro_catchment_name(self):
        self._create_micro_catchment("Existing catchment")

        with self.assertRaises(ValidationError) as context:
            self._create_micro_catchment("  existing CATCHMENT  ")

        self.assertEqual(
            context.exception.messages,
            ["Micro-catchment name 'existing CATCHMENT' already exists."],
        )

    def test_same_micro_catchment_name_is_allowed_in_different_districts(self):
        self._create_micro_catchment("Shared catchment")
        other_district = self._location("R", "DIST2")
        other_ta = self._location("D", "5678", other_district)
        other_gvh = self._location("W", "GVH-DIST2", other_ta)

        created = MicroCatchmentService(self.user).update_or_create(
            {
                "name": "shared CATCHMENT",
                "district_id": other_district.id,
                "ta_ids": [other_ta.id],
                "gvh_ids": [other_gvh.id],
                "audit_user_id": -1,
            }
        )

        self.assertEqual(created.district, other_district)
        self.assertEqual(created.code, "567801")

    def test_import_name_check_is_scoped_to_selected_district(self):
        self._create_micro_catchment("Shared import catchment")
        other_district = self._location("R", "DIST3")
        other_ta = self._location("D", "6789", other_district)
        other_gvh = self._location("W", "GVH-DIST3", other_ta)

        result = import_records(
            [
                {
                    "name": "shared IMPORT catchment",
                    "type": None,
                    "tas": [other_ta],
                    "gvhs": [other_gvh],
                }
            ],
            other_district,
            self.user,
        )

        self.assertEqual(result, {"created": 1, "updated": 0, "total": 1})
        self.assertTrue(
            MicroCatchment.objects.filter(
                district=other_district,
                name="shared IMPORT catchment",
            ).exists()
        )

    def test_import_reports_only_gvhs_assigned_to_another_catchment(self):
        self._create_micro_catchment("Existing catchment")
        available_gvh = self._location("W", "GVH2", self.ta)

        with self.assertRaises(ValidationError) as context:
            import_records(
                [
                    {
                        "name": "New catchment",
                        "type": None,
                        "tas": [self.ta],
                        "gvhs": [self.gvh, available_gvh],
                    }
                ],
                self.district,
                self.user,
            )

        message = " ".join(context.exception.messages)
        self.assertIn("GVH1 - GVH1", message)
        self.assertNotIn("GVH2", message)
