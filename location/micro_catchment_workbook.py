import csv
from io import BytesIO, StringIO

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.utils import TimeUtils

from .models import Location, MicroCatchment, MicroCatchmentGVH
from .services import MicroCatchmentService


CSV_HEADERS = (
    "district_code",
    "district_name",
    "ta_code",
    "ta_name",
    "gvh_code",
    "gvh_name",
    "micro_catchment_name",
)


HEADERS = (
    "name",
    "type",
    "district_code",
    "ta_codes",
    "gvh_codes",
)

EXPORT_HEADERS = (
    "code",
    "micro_catchment_name",
    "type",
    "district_code",
    "ta_code",
    "ta_name",
    "gvh_code",
    "gvh_name",
)


def build_template_workbook(district):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MicroCatchments"
    sheet.append(CSV_HEADERS)
    traditional_authorities = Location.objects.filter(
        type="D",
        parent=district,
        validity_to__isnull=True,
    ).order_by("code")
    rows_written = False
    for ta in traditional_authorities:
        gvhs = Location.objects.filter(
            type="W",
            parent=ta,
            validity_to__isnull=True,
        ).order_by("code")
        ta_has_gvhs = False
        for gvh in gvhs:
            sheet.append((district.code, district.name, ta.code, ta.name, gvh.code, gvh.name, ""))
            rows_written = ta_has_gvhs = True
        if not ta_has_gvhs:
            sheet.append((district.code, district.name, ta.code, ta.name, "", "", ""))
            rows_written = True
    if not rows_written:
        sheet.append((district.code, district.name, "", "", "", "", ""))
    _format_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _codes(value):
    return [code.strip() for code in _cell_text(value).split(",") if code.strip()]


def _format_sheet(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00677F")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max(len(_cell_text(cell.value)) for cell in column) + 2
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width, 12), 45)


def build_workbook(district):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MicroCatchments"
    sheet.append(EXPORT_HEADERS)

    catchments = (
        MicroCatchment.objects.filter(district=district, validity_to__isnull=True)
        .prefetch_related("traditional_authorities__location", "gvhs__location")
        .order_by("code")
    )
    for catchment in catchments:
        ta_locations = [
            link.location
            for link in catchment.traditional_authorities.filter(
                validity_to__isnull=True
            ).select_related("location")
        ]
        gvh_locations = [
            link.location
            for link in catchment.gvhs.filter(
                validity_to__isnull=True
            ).select_related("location__parent")
        ]

        rows = [
            (
                gvh.parent.code if gvh.parent else "",
                gvh.parent.name if gvh.parent else "",
                gvh.code,
                gvh.name,
            )
            for gvh in gvh_locations
        ]
        if not rows:
            rows = [(ta.code, ta.name, "", "") for ta in ta_locations] or [("", "", "", "")]

        for ta_code, ta_name, gvh_code, gvh_name in rows:
            sheet.append(
                (
                    catchment.code,
                    catchment.name,
                    catchment.type or "",
                    district.code,
                    ta_code,
                    ta_name,
                    gvh_code,
                    gvh_name,
                )
            )

    # Keep an immediately usable blank row when a district has no catchments yet.
    if sheet.max_row == 1:
        sheet.append(("", "", "", district.code, "", "", "", ""))

    _format_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_workbook(uploaded_file, district):
    try:
        workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
    except Exception as exc:
        raise ValidationError("The uploaded file is not a valid .xlsx workbook.") from exc

    if "MicroCatchments" not in workbook.sheetnames:
        raise ValidationError("The workbook must contain a 'MicroCatchments' sheet.")

    sheet = workbook["MicroCatchments"]
    rows = sheet.iter_rows(values_only=True)
    try:
        supplied_headers = tuple(_cell_text(value).lower() for value in next(rows))
    except StopIteration as exc:
        raise ValidationError("The MicroCatchments sheet is empty.") from exc

    missing = [header for header in HEADERS if header not in supplied_headers]
    if missing:
        raise ValidationError("Missing required columns: %s." % ", ".join(missing))
    indexes = {header: supplied_headers.index(header) for header in HEADERS}

    parsed = []
    seen_names = set()
    errors = []
    for row_number, row in enumerate(rows, start=2):
        values = {header: row[index] if index < len(row) else None for header, index in indexes.items()}
        if not any(_cell_text(value) for value in values.values()):
            continue
        name = _cell_text(values["name"])
        district_code = _cell_text(values["district_code"])
        row_errors = []
        if not name:
            row_errors.append("name is required")
        if district_code != district.code:
            row_errors.append(f"district_code must be {district.code}")
        if name.casefold() in seen_names:
            row_errors.append(f"name {name} appears more than once")
        seen_names.add(name.casefold())

        ta_codes = _codes(values["ta_codes"])
        gvh_codes = _codes(values["gvh_codes"])
        tas = list(
            Location.objects.filter(
                code__in=ta_codes,
                type="W",
                parent=district,
                validity_to__isnull=True,
            )
        )
        gvhs = list(
            Location.objects.filter(
                code__in=gvh_codes,
                type="V",
                parent__in=tas,
                validity_to__isnull=True,
            )
        )
        missing_tas = sorted(set(ta_codes) - {location.code for location in tas})
        missing_gvhs = sorted(set(gvh_codes) - {location.code for location in gvhs})
        if missing_tas:
            row_errors.append("invalid TA code(s): %s" % ", ".join(missing_tas))
        if missing_gvhs:
            row_errors.append("invalid GVH code(s): %s" % ", ".join(missing_gvhs))
        if not ta_codes:
            row_errors.append("at least one TA code is required")
        if not gvh_codes:
            row_errors.append("at least one GVH code is required")

        if row_errors:
            errors.append("Row %s: %s." % (row_number, "; ".join(row_errors)))
        else:
            parsed.append(
                {
                    "name": name,
                    "type": _cell_text(values["type"]) or None,
                    "tas": tas,
                    "gvhs": gvhs,
                }
            )

    if errors:
        raise ValidationError(errors)
    if not parsed:
        raise ValidationError("The workbook contains no completed micro-catchment rows.")
    return parsed


@transaction.atomic
def import_workbook(uploaded_file, district, user):
    records = parse_workbook(uploaded_file, district)
    return import_records(records, district, user)


def import_csv(uploaded_file, district, user):
    try:
        content = uploaded_file.read().decode("utf-8-sig")
        reader = csv.DictReader(StringIO(content))
        supplied_headers = tuple((header or "").strip().lower() for header in (reader.fieldnames or ()))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise ValidationError("The uploaded file is not a valid UTF-8 CSV file.") from exc

    missing = [header for header in CSV_HEADERS if header not in supplied_headers]
    if missing:
        raise ValidationError("Missing required columns: %s." % ", ".join(missing))

    grouped = {}
    errors = []
    for row_number, row in enumerate(reader, start=2):
        values = {(key or "").strip().lower(): _cell_text(value) for key, value in row.items()}
        name = values["micro_catchment_name"]
        ta_code = values["ta_code"]
        gvh_code = values["gvh_code"]
        if not name:
            continue

        row_errors = []
        if values["district_code"] != district.code:
            row_errors.append(f"district_code must be {district.code}")
        if not name:
            row_errors.append("micro_catchment_name is required")
        ta = Location.objects.filter(
            code=ta_code,
            type="D",
            parent=district,
            validity_to__isnull=True,
        ).first()
        if not ta:
            row_errors.append(f"invalid TA code: {ta_code}")
        gvh = None
        if gvh_code and ta:
            gvh = Location.objects.filter(
                code=gvh_code,
                type="W",
                parent=ta,
                validity_to__isnull=True,
            ).first()
            if not gvh:
                row_errors.append(f"invalid GVH code: {gvh_code}")
        if row_errors:
            errors.append("Row %s: %s." % (row_number, "; ".join(row_errors)))
            continue

        record = grouped.setdefault(
            name.casefold(),
            {"name": name, "tas": [], "gvhs": []},
        )
        if record["name"] != name:
            errors.append(f"Row {row_number}: micro_catchment_name is inconsistent across matching rows.")
        elif ta.id not in {location.id for location in record["tas"]}:
            record["tas"].append(ta)
        if gvh and gvh.id not in {location.id for location in record["gvhs"]}:
            record["gvhs"].append(gvh)

    if errors:
        raise ValidationError(errors)
    if not grouped:
        raise ValidationError("The CSV contains no completed micro-catchment rows.")

    records = [
        {
            "name": values["name"],
            "type": None,
            "tas": values["tas"],
            "gvhs": values["gvhs"],
        }
        for values in grouped.values()
    ]
    return import_records(records, district, user)


def import_excel(uploaded_file, district, user):
    try:
        workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
        sheet = workbook["MicroCatchments"]
        first_row = next(sheet.iter_rows(values_only=True))
        headers = tuple(_cell_text(value).lower() for value in first_row)
    except Exception as exc:
        raise ValidationError("The uploaded file is not a valid MicroCatchments workbook.") from exc

    uploaded_file.seek(0)
    if all(header in headers for header in CSV_HEADERS):
        output = StringIO(newline="")
        writer = csv.writer(output)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(row)
        converted = BytesIO(output.getvalue().encode("utf-8-sig"))
        return import_csv(converted, district, user)
    return import_workbook(uploaded_file, district, user)


@transaction.atomic
def import_records(records, district, user):
    validation_errors = []
    for record in records:
        name = record["name"].strip()
        if MicroCatchment.objects.filter(
            name__iexact=name,
            district=district,
            validity_to__isnull=True,
        ).exists():
            validation_errors.append(
                f"Micro-catchment name '{name}' already exists."
            )

        gvh_ids = [location.id for location in record["gvhs"]]
        conflicting_gvh_ids = MicroCatchmentGVH.objects.filter(
            location_id__in=gvh_ids,
            validity_to__isnull=True,
            micro_catchment__validity_to__isnull=True,
        ).values_list("location_id", flat=True)
        conflicting_gvhs = Location.objects.filter(
            id__in=conflicting_gvh_ids,
        ).order_by("code")
        if conflicting_gvhs.exists():
            conflict_labels = ", ".join(
                f"{gvh.code} - {gvh.name}" for gvh in conflicting_gvhs
            )
            validation_errors.append(
                f"Micro-catchment '{name}' contains GVH(s) already assigned to another "
                f"micro-catchment: {conflict_labels}."
            )

    if validation_errors:
        raise ValidationError(validation_errors)

    created = 0
    service = MicroCatchmentService(user)
    audit_user_id = user.id_for_audit

    for record in records:
        tas = record["tas"]
        gvhs = record["gvhs"]
        service.update_or_create(
            {
                **{
                    key: value
                    for key, value in record.items()
                    if key not in ("tas", "gvhs")
                },
                "district_id": district.id,
                "ta_ids": [location.id for location in tas],
                "gvh_ids": [location.id for location in gvhs],
                "audit_user_id": audit_user_id,
                "validity_from": TimeUtils.now(),
            }
        )
        created += 1

    return {"created": created, "updated": 0, "total": len(records)}
